"""Issue #5: isolated synthetic fixtures; --check-only never overwrites inputs."""
import argparse
import io
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = dict(cn="宋体", en="Times New Roman", size=12.0, spacing=1.5, left=3.0)
TEXT = ["完全自建合成正文第一段。", "完全自建合成正文第二段。", "合成第三段 ABC 0123456789。"]
SAMPLES = {
    "baseline_correct": {},
    "font_family_cn_wrong": {"cn": "楷体"},
    "font_ascii_wrong": {"en": "Arial"},
    "font_size_wrong": {"size": 10.5},
    "line_spacing_wrong": {"spacing": 1.0},
    "margin_wrong": {"left": 2.0},
    "style_override_control": {"inherit": True},
    "style_override_wrong": {"inherit": True, "override": True},
    "multisection_control": {"multi": True},
    "multisection_margin_wrong": {"multi": True, "second_left": 2.5},
}


def fonts(element, cn, en):
    element.font.name = en
    rpr = (element.element if hasattr(element, "element") else element._element).get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    for attr in list(rf.attrib):
        if attr.endswith("Theme"):
            del rf.attrib[attr]
    rf.set(qn("w:eastAsia"), cn)


def margins(section, left):
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(left)
    section.right_margin = Cm(2.5)


def build(config):
    doc = Document()
    normal = doc.styles["Normal"]
    fonts(normal, BASE["cn"], BASE["en"])
    normal.font.size = Pt(BASE["size"])
    normal.paragraph_format.line_spacing = BASE["spacing"]
    margins(doc.sections[0], config.get("left", BASE["left"]))

    def heading():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = BASE["spacing"]
        r = p.add_run("合成格式样本")
        fonts(r, "黑体", "黑体")
        r.font.size = Pt(16)
        r.bold = False

    def body(text, index):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = config.get("spacing", BASE["spacing"])
        r = p.add_run(text)
        if not config.get("inherit"):
            fonts(r, config.get("cn", BASE["cn"]), config.get("en", BASE["en"]))
            r.font.size = Pt(config.get("size", BASE["size"]))
        elif config.get("override") and index == 1:
            # Only this eastAsia attribute overrides Normal; size/ascii still inherit.
            rf = r._element.get_or_add_rPr().makeelement(qn("w:rFonts"), {})
            rf.set(qn("w:eastAsia"), "楷体")
            r._element.rPr.append(rf)

    heading()
    for i, text in enumerate(TEXT):
        body(text, i)
    if config.get("multi"):
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        margins(section, config.get("second_left", BASE["left"]))
        heading()
        for i, text in enumerate(TEXT[:2]):
            body(text, i)
    cp = doc.core_properties
    cp.author = "Synthetic fixture generator"
    cp.last_modified_by = ""
    cp.created = cp.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    return doc


def generate(directory):
    directory.mkdir(parents=True, exist_ok=True)
    for name, config in SAMPLES.items():
        stream = io.BytesIO()
        build(config).save(stream)
        with ZipFile(stream) as source, ZipFile(directory / (name + ".docx"), "w") as target:
            for part in sorted(source.namelist()):
                item = ZipInfo(part, date_time=(2000, 1, 1, 0, 0, 0))
                item.compress_type = ZIP_DEFLATED
                target.writestr(item, source.read(part))


def check(directory):
    """Validate all expected dimensions of the checked-in files, including inheritance."""
    for name, config in SAMPLES.items():
        path = directory / (name + ".docx")
        with ZipFile(path) as z:
            assert z.testzip() is None, name
            assert len(z.namelist()) == len(set(z.namelist())), name
        doc = Document(path)
        normal = doc.styles["Normal"]
        nr = normal.element.rPr
        rf = nr.find(qn("w:rFonts"))
        assert rf.get(qn("w:eastAsia")) == BASE["cn"], name
        assert rf.get(qn("w:ascii")) == BASE["en"], name
        assert normal.font.size.pt == BASE["size"], name
        assert not any(a.endswith("Theme") for a in rf.attrib), name
        assert len(doc.sections) == (2 if config.get("multi") else 1), name
        for i, sec in enumerate(doc.sections):
            expected_left = config.get("left", BASE["left"]) if i == 0 else config.get("second_left", BASE["left"])
            for side, cm in dict(top=2.5, bottom=2.5, right=2.5, left=expected_left).items():
                assert abs(getattr(sec, side + "_margin").cm - cm) < .01, (name, i, side)
        body = [p for p in doc.paragraphs if p.runs and p.alignment != WD_ALIGN_PARAGRAPH.CENTER]
        assert [p.text for p in body] == (TEXT + TEXT[:2] if config.get("multi") else TEXT), name
        for i, p in enumerate(body):
            assert p.paragraph_format.line_spacing == config.get("spacing", BASE["spacing"]), name
            assert len(p.runs) == 1, name
            r = p.runs[0]
            rpr = r._element.find(qn("w:rPr"))
            if config.get("inherit") and not (config.get("override") and i == 1):
                assert rpr is None, (name, "control must inherit")
            local_fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
            expected_cn = "楷体" if config.get("override") and i == 1 else config.get("cn", BASE["cn"])
            effective_cn = local_fonts.get(qn("w:eastAsia"), BASE["cn"]) if local_fonts is not None else BASE["cn"]
            effective_en = r.font.name or BASE["en"]
            effective_size = r.font.size.pt if r.font.size is not None else normal.font.size.pt
            assert (effective_cn, effective_en, effective_size) == (expected_cn, config.get("en", BASE["en"]), config.get("size", BASE["size"])), (name, i)
            if config.get("override") and i == 1:
                assert dict(local_fonts.attrib) == {qn("w:eastAsia"): "楷体"}, name
                assert len(rpr) == 1, name
        for p in doc.paragraphs:
            if p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
                assert p.text == "合成格式样本", name
                assert p.runs[0].font.size.pt == 16, name
                assert p.runs[0].font.bold is False, name
    print("PASS: 10 committed fixtures, package integrity, all baseline/error dimensions and inheritance")


def main():
    parser = argparse.ArgumentParser()
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--check", action="store_true", help="generate then validate")
    mode.add_argument("--check-only", action="store_true", help="validate existing files without rewriting")
    args = parser.parse_args()
    directory = Path(__file__).resolve().parent
    if not args.check_only:
        generate(directory)
    if args.check or args.check_only:
        check(directory)


if __name__ == "__main__":
    main()
