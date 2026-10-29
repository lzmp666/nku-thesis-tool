"""
生成两份 synthetic .docx：
1. clean_sample.docx        无格式问题
2. known_issues_sample.docx 含三项已知问题：东亚字体、西文字体、字号

所有内容为合成文本，可重复生成，不涉及真实论文或身份信息。
规则 ID 使用 demo/* 前缀，仅用于演示/测试。
"""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


OUT = Path(__file__).parent


def set_east_asian_font(
    run, font_name: str, western_font: str = "Times New Roman"
) -> None:
    """设置 run 的东亚字体和对应的西文字体。"""
    run.font.name = western_font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = r_pr.makeelement(qn("w:rFonts"), {})
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), font_name)


def build_clean(path: Path) -> None:
    """生成无格式问题的样本。"""
    doc = Document()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("这是一段合成演示文本，用于验证格式诊断链路。")
    set_east_asian_font(run, "宋体", "Times New Roman")
    run.font.size = Pt(12)
    doc.save(path)


def build_known_issues(path: Path) -> None:
    """生成带有正文边界且含三项已知问题的样本。"""
    doc = Document()
    doc.add_heading("第一章 绪论", level=1)

    # 问题 1：东亚字体应为宋体，这里故意用黑体
    paragraph1 = doc.add_paragraph()
    run1 = paragraph1.add_run("【问题1】东亚字体不符合预期。")
    set_east_asian_font(run1, "黑体", "Times New Roman")
    run1.font.size = Pt(12)

    # 问题 2：西文字体应为 Times New Roman，这里故意用 Arial
    paragraph2 = doc.add_paragraph()
    run2 = paragraph2.add_run("【问题2】Latin 西文字体不符合预期。")
    set_east_asian_font(run2, "宋体", "Arial")
    run2.font.size = Pt(12)

    # 问题 3：字号应为 12pt，这里故意用 10.5pt
    paragraph3 = doc.add_paragraph()
    run3 = paragraph3.add_run("【问题3】字号不符合预期。")
    set_east_asian_font(run3, "宋体", "Times New Roman")
    run3.font.size = Pt(10.5)

    doc.add_heading("参考文献", level=1)
    doc.save(path)


if __name__ == "__main__":
    build_clean(OUT / "clean_sample.docx")
    build_known_issues(OUT / "known_issues_sample.docx")
    print("generated:", OUT)
